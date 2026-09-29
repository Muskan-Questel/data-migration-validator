from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from collections import Counter
import unicodedata
import re

import pandas as pd
from openpyxl.styles import PatternFill


@dataclass(frozen=True)
class ComparisonResult:
    sheets: dict[str, pd.DataFrame]
    summary: dict[str, Any]
    summary_rows: list[dict[str, Any]]
    highlighted_cells: dict[str, dict[tuple[int, int], str]]
    key_columns: list[str]


def compare_dataframes(
    excel_df: pd.DataFrame,
    db_df: pd.DataFrame,
    key_columns: list[str],
    optional_key_columns: list[str] | None = None,
    column_aliases: dict[str, str] | None = None,
    header_aliases: dict[str, list[str]] | None = None,
    summary_subject: str = "Records",
) -> ComparisonResult:
    """Compare uploaded Excel data to a DB extract using Excel columns as contract."""
    aliases = column_aliases or {}
    excel_df = excel_df.copy()
    original_db_columns = list(db_df.columns)
    db_df = db_df.rename(columns=aliases).copy()
    excel_header_mapping = _excel_header_mapping(list(excel_df.columns), header_aliases or {})
    excel_df = excel_df.rename(columns=excel_header_mapping)
    excel_df = _normalize_party_columns(excel_df)
    db_df = _normalize_party_columns(db_df)

    excel_columns = list(excel_df.columns)
    db_columns = list(db_df.columns)
    resolved_key_columns = _resolve_key_columns(
        key_columns,
        optional_key_columns or [],
        excel_columns,
        db_columns,
    )

    missing_keys = [
        key
        for key in key_columns
        if key not in excel_columns or key not in db_columns
    ]
    if missing_keys:
        raise ValueError(
            "Configured key columns must exist in both Excel and DB extract: "
            + ", ".join(missing_keys)
        )
    if not resolved_key_columns:
        raise ValueError("At least one configured key column must exist in both Excel and DB extract.")

    missing_db_columns = [column for column in excel_columns if column not in db_columns]
    compare_columns = [column for column in excel_columns if column not in missing_db_columns]
    duplicate_records, excel_dup_mask, db_dup_mask = _duplicate_records(
        excel_df,
        db_df,
        resolved_key_columns,
    )
    scientific_warnings, scientific_highlights = _scientific_number_warnings(
        excel_df,
        db_df,
        resolved_key_columns,
        compare_columns,
        excel_df.attrs.get("scientific_number_values", []),
    )

    excel_unique = excel_df.loc[~excel_dup_mask].copy()
    db_unique = db_df.loc[~db_dup_mask].copy()

    excel_by_key = _rows_by_key(excel_unique, resolved_key_columns)
    db_by_key = _rows_by_key(db_unique, resolved_key_columns)

    excel_keys = set(excel_by_key)
    db_keys = set(db_by_key)
    matched_keys = sorted(excel_keys & db_keys, key=str)

    extra_in_excel = _keys_for_records(excel_by_key, excel_keys - db_keys, resolved_key_columns)
    extra_in_database = _keys_for_records(db_by_key, db_keys - excel_keys, resolved_key_columns)
    mismatched_records, highlighted_cells = _mismatched_records(
        excel_by_key,
        db_by_key,
        matched_keys,
        resolved_key_columns,
        compare_columns,
        excel_df.attrs.get("scientific_number_values", []),
    )

    summary = {
        "excel_rows": len(excel_df),
        "db_rows": len(db_df),
        "compared_columns": len(compare_columns),
        "matched_rows": len(matched_keys),
        "mismatched_records": mismatched_records[
            resolved_key_columns
        ].drop_duplicates().shape[0]
        if not mismatched_records.empty
        else 0,
        "extra_records_in_database": len(extra_in_database),
        "extra_records_in_excel": len(extra_in_excel),
        "duplicate_excel_rows": int(excel_dup_mask.sum()),
        "duplicate_db_rows": int(db_dup_mask.sum()),
        "missing_db_columns": len(missing_db_columns),
        "scientific_number_warnings": len(scientific_warnings),
    }
    summary_rows = _summary_rows(summary, summary_subject)
    summary_sheet = pd.DataFrame(summary_rows, columns=["Metric", "Value"])
    column_mapping_sheet = _column_mapping_sheet(
        excel_columns,
        original_db_columns,
        aliases,
        excel_header_mapping,
    )

    return ComparisonResult(
        sheets={
            "Summary": summary_sheet,
            "Column_Mapping": column_mapping_sheet,
            "Mismatched_Records": mismatched_records,
            "Only in Database": extra_in_database,
            "Only in Excel": extra_in_excel,
            "Duplicate Records": duplicate_records,
            "Scientific_Number_Warnings": scientific_warnings,
        },
        summary=summary,
        summary_rows=summary_rows,
        key_columns=resolved_key_columns,
        highlighted_cells={
            "Mismatched_Records": highlighted_cells,
            "Scientific_Number_Warnings": scientific_highlights,
        },
    )


def _resolve_key_columns(
    required_key_columns: list[str],
    optional_key_columns: list[str],
    excel_columns: list[str],
    db_columns: list[str],
) -> list[str]:
    resolved_key_columns = list(required_key_columns)
    for column in optional_key_columns:
        if column in resolved_key_columns:
            continue
        if column in excel_columns and column in db_columns:
            resolved_key_columns.append(column)
    return resolved_key_columns


def _duplicate_records(
    excel_df: pd.DataFrame,
    db_df: pd.DataFrame,
    key_columns: list[str],
) -> tuple[pd.DataFrame, pd.Series, pd.Series]:
    excel_keys = [
        tuple(_normalize_key_value(row[column]) for column in key_columns)
        for _, row in excel_df.iterrows()
    ]
    db_keys = [
        tuple(_normalize_key_value(row[column]) for column in key_columns)
        for _, row in db_df.iterrows()
    ]
    excel_counts = Counter(excel_keys)
    db_counts = Counter(db_keys)
    excel_dup_mask = pd.Series(
        [excel_counts[key] > 1 for key in excel_keys], index=excel_df.index
    )
    db_dup_mask = pd.Series(
        [db_counts[key] > 1 for key in db_keys], index=db_df.index
    )

    rows: list[dict[str, Any]] = []
    output_columns = [
        "Source",
        "Duplicate Group",
        "Duplicate Count",
        "Source Row",
        *key_columns,
        *[f"Excel: {column}" for column in excel_df.columns if column not in key_columns],
        *[f"Database: {column}" for column in db_df.columns if column not in key_columns],
    ]
    for source, df, keys, counts, mask in (
        ("Excel", excel_df, excel_keys, excel_counts, excel_dup_mask),
        ("Database", db_df, db_keys, db_counts, db_dup_mask),
    ):
        group_ids: dict[tuple[Any, ...], str] = {}
        group_number = 0
        for position, (_, row) in enumerate(df.iterrows()):
            key = keys[position]
            if not bool(mask.iloc[position]):
                continue
            if key not in group_ids:
                group_number += 1
                group_ids[key] = f"{source[:3].upper()}-{group_number:03d}"
            record: dict[str, Any] = {
                "Source": source,
                "Duplicate Group": group_ids[key],
                "Duplicate Count": counts[key],
                "Source Row": position + 2,
                **{column: row[column] for column in key_columns},
            }
            prefix = "Excel: " if source == "Excel" else "Database: "
            record.update(
                {
                    f"{prefix}{column}": row[column]
                    for column in df.columns
                    if column not in key_columns
                }
            )
            rows.append(record)
    return (
        pd.DataFrame(rows, columns=output_columns),
        excel_dup_mask,
        db_dup_mask,
    )


def write_report(result: ComparisonResult, output_path: str) -> None:
    mismatch_fills = {
        "mismatch": PatternFill(
            fill_type="solid",
            start_color="FFF2CC",
            end_color="FFF2CC",
        ),
        "whitespace": PatternFill(
            fill_type="solid",
            start_color="FFC6EFCE",
            end_color="FFC6EFCE",
        ),
        "scientific": PatternFill(
            fill_type="solid",
            start_color="FFF2CC",
            end_color="FFF2CC",
        ),
        "scientific_both": PatternFill(
            fill_type="solid",
            start_color="FFBDD7EE",
            end_color="FFBDD7EE",
        ),
    }
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, sheet_df in result.sheets.items():
            sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.book[sheet_name]
            worksheet.freeze_panes = "A2"
            for column_cells in worksheet.columns:
                max_length = max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in column_cells
                )
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                    max(max_length + 2, 12),
                    45,
                )
        for sheet_name, cells in result.highlighted_cells.items():
            worksheet = writer.book[sheet_name]
            for (row_number, column_number), highlight_type in cells.items():
                worksheet.cell(row=row_number, column=column_number).fill = mismatch_fills[
                    highlight_type
                ]


def _mismatched_records(
    excel_by_key: dict[tuple[Any, ...], pd.Series],
    db_by_key: dict[tuple[Any, ...], pd.Series],
    matched_keys: list[tuple[Any, ...]],
    key_columns: list[str],
    compare_columns: list[str],
    raw_values: list[tuple[int, str, Any]],
) -> tuple[pd.DataFrame, dict[tuple[int, int], str]]:
    raw_cells = {
        (row_index, _normalize_header(column)): value
        for row_index, column, value in raw_values
    }
    mismatch_rows: list[tuple[pd.Series, pd.Series, list[str]]] = []
    included_value_columns: set[str] = set()
    for key in matched_keys:
        excel_row = excel_by_key[key]
        db_row = db_by_key[key]
        mismatched_columns = [
            column
            for column in compare_columns
            if column not in key_columns
            and (
                not _comparison_equal(excel_row[column], db_row[column], column)
                or _formatting_only_difference(excel_row[column], db_row[column])
            )
        ]
        if mismatched_columns:
            mismatch_rows.append((excel_row, db_row, mismatched_columns))
            included_value_columns.update(mismatched_columns)

    value_columns = [
        column
        for column in compare_columns
        if column not in key_columns and column in included_value_columns
    ]
    columns = [
        *key_columns,
        *[
            source_column
            for column in value_columns
            for source_column in (f"{column} (Excel)", f"{column} (Database)")
        ],
    ]
    column_numbers = {column: index + 1 for index, column in enumerate(columns)}
    rows: list[dict[str, Any]] = []
    highlighted_cells: dict[tuple[int, int], str] = {}
    for excel_row, db_row, mismatched_columns in mismatch_rows:
        output_row = {
            **{column: excel_row[column] for column in key_columns},
            **{
                source_column: value
                for column in value_columns
                for source_column, value in (
                    (
                        f"{column} (Excel)",
                        raw_cells.get(
                            (excel_row.name, _normalize_header(column)),
                            excel_row[column],
                        ),
                    ),
                    (f"{column} (Database)", db_row[column]),
                )
            },
        }
        rows.append(output_row)

        sheet_row = len(rows) + 1
        for column in mismatched_columns:
            highlight_type = (
                "whitespace"
                if _formatting_only_difference(excel_row[column], db_row[column])
                else "mismatch"
            )
            highlighted_cells[(sheet_row, column_numbers[f"{column} (Excel)"])] = highlight_type
            highlighted_cells[(sheet_row, column_numbers[f"{column} (Database)"])] = highlight_type

    return pd.DataFrame(rows, columns=columns), highlighted_cells


_SCIENTIFIC_NUMBER = re.compile(
    r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)[eE][+-]?\d+$"
)


def _scientific_number_warnings(
    excel_df: pd.DataFrame,
    db_df: pd.DataFrame,
    key_columns: list[str],
    compare_columns: list[str],
    raw_values: list[tuple[int, str, Any]],
) -> tuple[pd.DataFrame, dict[tuple[int, int], str]]:
    rows: list[dict[str, Any]] = []
    highlighted_cells: dict[tuple[int, int], str] = {}
    raw_cells = {
        (row_index, _normalize_header(column)): value
        for row_index, column, value in raw_values
    }
    excel_by_key = _rows_by_key(excel_df, key_columns)
    db_by_key = _rows_by_key(db_df, key_columns)
    for key in sorted(set(excel_by_key) & set(db_by_key), key=str):
        excel_row = excel_by_key[key]
        db_row = db_by_key[key]
        key_values = {column: excel_row[column] for column in key_columns}
        for column in compare_columns:
            excel_value = raw_cells.get(
                (excel_row.name, _normalize_header(column)),
                excel_row[column],
            )
            db_value = db_row[column]
            excel_scientific = _is_scientific_number(excel_value)
            db_scientific = _is_scientific_number(db_value)
            if not excel_scientific and not db_scientific:
                continue
            status = (
                "Both Excel and Database scientific"
                if excel_scientific and db_scientific
                else "Excel scientific"
                if excel_scientific
                else "Database scientific"
            )
            rows.append(
                {
                    **key_values,
                    "Field": column,
                    "Excel Value": excel_value,
                    "Database Value": db_value,
                    "Status": status,
                }
            )
            highlight_type = "scientific_both" if excel_scientific and db_scientific else "scientific"
            sheet_row = len(rows) + 1
            field_column = len(key_columns) + 1
            highlighted_cells[(sheet_row, field_column)] = highlight_type
            highlighted_cells[(sheet_row, field_column + 1)] = highlight_type
            highlighted_cells[(sheet_row, field_column + 2)] = highlight_type
    return pd.DataFrame(
        rows,
        columns=[*key_columns, "Field", "Excel Value", "Database Value", "Status"],
    ), highlighted_cells


def _is_scientific_number(value: Any) -> bool:
    if pd.isna(value):
        return False
    return bool(_SCIENTIFIC_NUMBER.fullmatch(str(value).strip()))


def _keys_for_records(
    rows_by_key: dict[tuple[Any, ...], pd.Series],
    keys: set[tuple[Any, ...]],
    key_columns: list[str],
) -> pd.DataFrame:
    rows = [
        dict(zip(key_columns, key, strict=True))
        for key in sorted(keys, key=str)
    ]
    return pd.DataFrame(rows, columns=key_columns)


def _column_mapping_sheet(
    excel_columns: list[str],
    db_columns: list[str],
    aliases: dict[str, str],
    excel_header_mapping: dict[str, str] | None = None,
) -> pd.DataFrame:
    canonical_by_uploaded_column = {
        canonical_column: uploaded_column
        for uploaded_column, canonical_column in (excel_header_mapping or {}).items()
    }
    db_column_by_excel_column = {
        aliases.get(db_column, db_column): db_column
        for db_column in db_columns
        if aliases.get(db_column, db_column) in excel_columns
    }
    mapped_db_columns = set(db_column_by_excel_column.values())
    rows = [
        {
            "Excel Column": canonical_by_uploaded_column.get(excel_column, excel_column),
            "DB Column": db_column_by_excel_column.get(excel_column, ""),
        }
        for excel_column in excel_columns
    ]
    rows.extend(
        {"Excel Column": "", "DB Column": db_column}
        for db_column in db_columns
        if db_column not in mapped_db_columns
    )
    return pd.DataFrame(rows, columns=["Excel Column", "DB Column"])


def _excel_header_mapping(
    excel_columns: list[str],
    header_aliases: dict[str, list[str]],
) -> dict[str, str]:
    alias_to_canonical = _alias_to_canonical_column(header_aliases)
    mapping: dict[str, str] = {}
    mapped_canonical_columns: dict[str, str] = {}

    for excel_column in excel_columns:
        literal_header = str(excel_column).replace("\u00a0", " ").strip().casefold()
        if literal_header in {"applicant(s)", "inventor(s)"}:
            continue
        canonical_column = alias_to_canonical.get(_normalize_header(excel_column))
        if not canonical_column or canonical_column == excel_column:
            continue
        if canonical_column in mapped_canonical_columns:
            raise ValueError(
                "Multiple uploaded headers map to the same configured header "
                f"{canonical_column}: {mapped_canonical_columns[canonical_column]}, "
                f"{excel_column}"
            )
        mapping[excel_column] = canonical_column
        mapped_canonical_columns[canonical_column] = excel_column

    return mapping


def _alias_to_canonical_column(header_aliases: dict[str, list[str]]) -> dict[str, str]:
    alias_to_canonical: dict[str, str] = {}
    for canonical_column, aliases in header_aliases.items():
        for alias in [canonical_column, *aliases]:
            normalized_alias = _normalize_header(alias)
            existing_column = alias_to_canonical.get(normalized_alias)
            if existing_column and existing_column != canonical_column:
                raise ValueError(
                    f"Header alias {alias!r} is configured for both "
                    f"{existing_column!r} and {canonical_column!r}"
                )
            alias_to_canonical[normalized_alias] = canonical_column
    return alias_to_canonical


def _normalize_header(value: Any) -> str:
    return "".join(
        character.lower()
        for character in str(value).replace("\u00a0", " ")
        if character.isalnum()
    )


def _exact_equal(left: Any, right: Any) -> bool:
    if _is_blank(left) and _is_blank(right):
        return True
    if pd.isna(left) and pd.isna(right):
        return True

    if isinstance(left, str) and isinstance(right, str):
        return left.casefold() == right.casefold()

    if type(left) is not type(right):
        return False
    return bool(left == right)


def _comparison_equal(left: Any, right: Any, column: str) -> bool:
    if _is_party_column(column):
        return _party_values_equal(left, right)
    if _mapped_value_difference(left, right, column):
        return True
    if _exact_equal(left, right):
        return True
    left_number = _numeric_value(left)
    right_number = _numeric_value(right)
    return left_number is not None and right_number is not None and left_number == right_number


def _normalize_party_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for column in df.columns:
        if _is_party_column(column):
            df[column] = df[column].map(_semicolon_list)
    return df


def _is_party_column(column: str) -> bool:
    normalized = _normalize_header(column)
    return normalized in {
        "applicant", "applicants", "inventor", "inventors", "associate", "associates"
    }
def _semicolon_list(value: Any) -> Any:
    if _is_blank(value) or pd.isna(value):
        return ""
    return ";".join(part.strip() for part in str(value).split(";") if part.strip())


def _party_values_equal(left: Any, right: Any) -> bool:
    return Counter(_party_parts(left)) == Counter(_party_parts(right))


def _party_parts(value: Any) -> list[str]:
    normalized = _semicolon_list(value)
    return [part for part in str(normalized).split(";") if part]


def _numeric_value(value: Any) -> float | None:
    if isinstance(value, bool) or _is_blank(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not text or (len(text) > 1 and text[0] == "0" and not text.startswith("0.")):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _is_blank(value: Any) -> bool:
    if value is None or (isinstance(value, str) and not value.strip()):
        return True
    if not pd.api.types.is_scalar(value):
        return False
    return bool(pd.isna(value))


def _whitespace_only_difference(left: Any, right: Any) -> bool:
    return (
        isinstance(left, str)
        and isinstance(right, str)
        and left != right
        and left.strip() == right.strip()
    )


def _formatting_only_difference(left: Any, right: Any) -> bool:
    return (
        isinstance(left, str)
        and isinstance(right, str)
        and left != right
        and left.strip().casefold() == right.strip().casefold()
    )


_ACCEPTED_VALUE_MAPPINGS = {
    "organisationcategory": {
        "client": "client",
        "associate": "foreign associate",
        "agent": "foreign associate",
        "other side client": "other side client",
        "otherside client": "other side client",
        "otherside associate": "other side solicitor",
        "other side associate": "other side solicitor",
        "otherside solicitor": "other side solicitor",
        "other side solicitor": "other side solicitor",
        "other associate": "other associate",
    },
}


def _mapped_value_difference(left: Any, right: Any, column: str) -> bool:
    mappings = {
        _normalize_header(key): _normalize_header(value)
        for key, value in _ACCEPTED_VALUE_MAPPINGS.get(
            _normalize_header(column), {}
        ).items()
    }
    if not mappings or not isinstance(left, str) or not isinstance(right, str):
        return False
    left_value = _normalize_header(left)
    right_value = _normalize_header(right)
    return mappings.get(left_value) == right_value or mappings.get(right_value) == left_value


def _normalize_key_value(value: Any) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    normalized = unicodedata.normalize("NFKC", str(value))
    normalized = normalized.translate(
        str.maketrans("", "", "\u200b\u200c\u200d\ufeff")
    )
    return " ".join(normalized.split()).casefold()


def _rows_by_key(df: pd.DataFrame, key_columns: list[str]) -> dict[tuple[Any, ...], pd.Series]:
    rows: dict[tuple[Any, ...], pd.Series] = {}
    for _, row in df.iterrows():
        key = tuple(_normalize_key_value(row[column]) for column in key_columns)
        rows[key] = row
    return rows


def _label(value: str) -> str:
    return value.replace("_", " ").title()


def _summary_rows(summary: dict[str, Any], subject: str) -> list[dict[str, Any]]:
    subject = subject.strip() or "Records"
    metric_labels = {
        "excel_rows": f"{subject} in Excel",
        "db_rows": f"{subject} in Database",
        "compared_columns": "Fields Compared",
        "matched_rows": f"{subject} Matched by Identifier",
        "mismatched_records": f"{subject} with Data Differences",
        "extra_records_in_database": f"{subject} Only in Database",
        "extra_records_in_excel": f"{subject} Only in Excel",
        "duplicate_excel_rows": f"Duplicate {subject} in Excel",
        "duplicate_db_rows": f"Duplicate {subject} in Database",
        "missing_db_columns": "Excel Columns Missing in Database",
        "scientific_number_warnings": "Scientific-Notation Warnings",
    }
    return [
        {"Metric": metric_labels.get(key, _label(key)), "Value": value}
        for key, value in summary.items()
    ]
